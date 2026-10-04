"""Тесты «простого» слоя: спецификации, заказы с отметками, дефицит, импорт с отступами."""
import io
from decimal import Decimal

import openpyxl

from .conftest import item_by_code


def _xlsx(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_with_indents_applies_immediately(client, admin):
    data = _xlsx([
        ["Обозначение", "Наименование", "Кол-во"],
        ["ИНД-1.00.000", "Изделие с отступами", 1],
        ["    ИНД-1.01.000", "Узел", 2],
        ["        ИНД-1.01.001", "Деталь узла", 3],
        ["    ВИНТ М6х20 ГОСТ 11738", "Винт", 4],
    ])
    r = client.post("/api/simple/import", headers=admin, files={"file": ("s.xlsx", data)})
    assert r.status_code == 200, r.text
    assert r.json()["applied"] is True and r.json()["item_id"]
    spec = client.get(f"/api/simple/specs/{r.json()['item_id']}", headers=admin).json()
    tree = spec["tree"]
    assert tree["code"] == "ИНД-1.00.000" and len(tree["children"]) == 2
    node = next(c for c in tree["children"] if c["code"] == "ИНД-1.01.000")
    assert node["children"][0]["code"] == "ИНД-1.01.001" and Decimal(node["children"][0]["qty"]) == 3
    assert next(c for c in tree["children"] if c["code"].startswith("ВИНТ"))["item_type"] == "fastener"


def test_import_without_hierarchy_asks_for_mapping(client, admin):
    data = _xlsx([["Артикул", "Название", "Кол"], ["A1", "a", 1], ["A2", "b", 2]])
    r = client.post("/api/simple/import", headers=admin, files={"file": ("f.xlsx", data)})
    assert r.status_code == 200 and r.json()["applied"] is False and r.json()["token"]


def test_spec_inline_editing_with_draft_and_apply(client, admin, product_id):
    spec = client.get(f"/api/simple/specs/{product_id}", headers=admin).json()
    before_rev = spec["rev"]
    # добавить новую строку
    r = client.post(f"/api/simple/specs/{product_id}/lines", headers=admin, json={"code": "ТАБЛ-1", "name": "Табличка", "qty": 1, "item_type": "purchased"})
    assert r.status_code == 200, r.text
    spec = client.get(f"/api/simple/specs/{product_id}", headers=admin).json()
    assert spec["has_draft"] is True
    line = next(c for c in spec["tree"]["children"] if c["code"] == "ТАБЛ-1")
    # изменить количество и переименовать
    r = client.patch(f"/api/simple/specs/{product_id}/lines/{line['line_id']}", headers=admin, json={"qty": 2, "name": "Табличка фирменная"})
    assert r.status_code == 200
    spec = client.get(f"/api/simple/specs/{product_id}", headers=admin).json()
    line = next(c for c in spec["tree"]["children"] if c["code"] == "ТАБЛ-1")
    assert Decimal(line["qty"]) == 2 and line["name"] == "Табличка фирменная"
    # применить — новая действующая редакция
    r = client.post(f"/api/simple/specs/{product_id}/apply", headers=admin)
    assert r.status_code == 200 and r.json()["applied"]
    spec = client.get(f"/api/simple/specs/{product_id}", headers=admin).json()
    assert spec["has_draft"] is False and spec["rev"] != before_rev
    # удалить и отменить правку
    line = next(c for c in spec["tree"]["children"] if c["code"] == "ТАБЛ-1")
    client.delete(f"/api/simple/specs/{product_id}/lines/{line['line_id']}", headers=admin)
    assert client.get(f"/api/simple/specs/{product_id}", headers=admin).json()["has_draft"] is True
    client.post(f"/api/simple/specs/{product_id}/discard", headers=admin)
    spec = client.get(f"/api/simple/specs/{product_id}", headers=admin).json()
    assert spec["has_draft"] is False and any(c["code"] == "ТАБЛ-1" for c in spec["tree"]["children"])


def test_order_checklist_and_shortage(client, admin, product_id):
    r = client.post("/api/simple/orders", headers=admin, json={"item_id": product_id, "qty": 2, "due_date": "2030-01-15", "customer": "Тест"})
    assert r.status_code == 200, r.text
    kid = r.json()["id"]
    o = client.get(f"/api/simple/orders/{kid}", headers=admin).json()
    assert o["percent"] == 0 and o["tree"]
    bearing = next(l for l in o["tree"] if l["item"]["code"] == "ПОДШ 6206")
    assert Decimal(bearing["required"]) == 8  # 4 × 2
    # частичное поступление
    r = client.post(f"/api/simple/orders/{kid}/lines/{bearing['id']}", headers=admin, json={"received": 3})
    assert r.json()["done"] is False
    sh = client.get("/api/simple/shortage", headers=admin).json()
    row = next(x for x in sh if x["item"]["code"] == "ПОДШ 6206")
    assert any(oo["id"] == kid and Decimal(oo["qty"]) == 5 for oo in row["orders"])
    # галочка — всё пришло
    r = client.post(f"/api/simple/orders/{kid}/lines/{bearing['id']}", headers=admin, json={"done": True})
    assert r.json()["done"] is True and Decimal(r.json()["received"]) == 8
    o = client.get(f"/api/simple/orders/{kid}", headers=admin).json()
    assert next(l for l in o["tree"] if l["id"] == bearing["id"])["status"] == "ready"
    # сборочный узел: галочка на узле закрывает его целиком
    body = next(l for l in o["tree"] if l["item"]["code"] == "РЧ-100.01.000")
    client.post(f"/api/simple/orders/{kid}/lines/{body['id']}", headers=admin, json={"done": True})
    o = client.get(f"/api/simple/orders/{kid}", headers=admin).json()
    assert next(l for l in o["tree"] if l["id"] == body["id"])["status"] == "ready"
    # список заказов и экспорт
    lst = client.get("/api/simple/orders", headers=admin).json()
    assert any(x["id"] == kid and x["status"] == "in_work" for x in lst)
    r = client.get(f"/api/simple/orders/{kid}/export.xlsx", headers=admin)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.active.max_row > 5
    r = client.get("/api/simple/shortage/export.xlsx", headers=admin)
    assert r.status_code == 200


def test_shortage_request_creates_draft_orders(client, admin):
    r = client.post("/api/simple/shortage/request", headers=admin)
    assert r.status_code == 200
    orders = r.json()["orders"]
    assert orders and all(o["lines"] > 0 for o in orders)
    po = client.get("/api/purchasing/orders", headers=admin, params={"status": "draft"}).json()
    assert any(p["comment"] == "Из ведомости дефицита" for p in po)


def test_specs_list(client, admin):
    specs = client.get("/api/simple/specs", headers=admin).json()
    assert any(s["code"] == "РЧ-100.00.000" and s["positions"] > 5 for s in specs)
    assert all("has_draft" in s and "open_orders" in s for s in specs)
    assert item_by_code(client, admin, "ПОДШ 6206")["code"] not in {s["code"] for s in specs}
