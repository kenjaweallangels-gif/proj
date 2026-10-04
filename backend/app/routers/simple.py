"""«Простой» слой API для основного интерфейса плановика.

Словарь интерфейса: Спецификация (= изделие с составом), Заказ (= комплект: изделие × кол-во × срок),
Поступление (= отметка «пришло» в заказе), Дефицит (= чего не хватает по открытым заказам).
Отметки в заказе НЕ двигают склад — это рабочий чек-лист плановика. Складской учёт — в разделе «Ещё».
"""
from __future__ import annotations

import hashlib
import io
from datetime import date
from decimal import Decimal

from fastapi import APIRouter, Depends, File, HTTPException, Response, UploadFile
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import audit
from ..db import get_db
from ..models import (
    BomLine,
    Item,
    ItemRevision,
    ItemType,
    Kit,
    KitLine,
    Lifecycle,
    PlanLine,
    POLine,
    PurchaseOrder,
    User,
)
from ..security import check_clearance, require
from ..services import bom as bom_svc
from ..services import excel_import
from ..services import planning as plan_svc
from .misc import _uploads

router = APIRouter(prefix="/api/simple", tags=["simple"])
Q4 = Decimal("0.0001")


def _item_brief(i: Item) -> dict:
    return {"id": i.id, "code": i.code, "name": i.name, "item_type": i.item_type.value, "unit": i.unit}


# ------------------------------------------------------------ спецификации --
@router.get("/specs")
def specs(db: Session = Depends(get_db), me: User = Depends(require("items:read"))):
    """Изделия верхнего уровня (ни во что не входят), с размером состава и числом открытых заказов."""
    used = {r[0] for r in db.execute(select(BomLine.child_item_id).distinct())}
    maxc = 99 if "*" in me.permissions else me.clearance
    open_kits: dict[int, int] = {}
    for k in db.query(Kit).filter(Kit.status.in_(["open", "in_work"])).all():
        open_kits[k.item_id] = open_kits.get(k.item_id, 0) + 1
    out = []
    for it in db.query(Item).filter(Item.confidentiality <= maxc).order_by(Item.code).all():
        rev = it.current_revision
        if it.id in used or not rev or not rev.lines:
            if it.item_type != ItemType.product or it.id in used:
                continue
        tree = bom_svc.build_tree(db, it, Decimal(1), 30, {}, {})
        nodes = bom_svc.flatten(tree)
        draft = any(r.status == Lifecycle.draft for r in it.revisions)
        out.append({**_item_brief(it), "rev": rev.rev if rev else None, "positions": len(nodes) - 1,
                    "has_draft": draft, "open_orders": open_kits.get(it.id, 0), "updated_at": it.updated_at})
    return out


class SpecNode(BaseModel):
    line_id: int | None
    item_id: int
    code: str
    name: str
    item_type: str
    unit: str
    qty: Decimal
    level: int
    children: list["SpecNode"] = []


def _edit_rev(it: Item) -> ItemRevision | None:
    """Редакция для редактора: черновик, если есть, иначе действующая."""
    return next((r for r in reversed(it.revisions) if r.status == Lifecycle.draft), None) or it.current_revision


def _spec_tree(it: Item, qty: Decimal = Decimal(1), level: int = 0, line_id: int | None = None, depth: int = 0) -> SpecNode:
    node = SpecNode(line_id=line_id, item_id=it.id, code=it.code, name=it.name, item_type=it.item_type.value, unit=it.unit, qty=qty, level=level)
    rev = _edit_rev(it)
    if rev and depth < 30:
        node.children = [_spec_tree(ln.child, ln.qty, level + 1, ln.id, depth + 1) for ln in rev.lines]
    return node


@router.get("/specs/{item_id}")
def spec(item_id: int, db: Session = Depends(get_db), me: User = Depends(require("items:read"))):
    it = db.get(Item, item_id)
    if not it:
        raise HTTPException(404, "Спецификация не найдена")
    check_clearance(me, it.confidentiality)
    rev = it.current_revision
    has_draft = False
    seen: set = set()

    def scan(node: Item):
        nonlocal has_draft
        if node.id in seen:
            return
        seen.add(node.id)
        if any(r.status == Lifecycle.draft for r in node.revisions):
            has_draft = True
        r = _edit_rev(node)
        for ln in r.lines if r else []:
            scan(ln.child)

    scan(it)
    released = [r for r in it.revisions if r.status == Lifecycle.released]
    return {"item": _item_brief(it), "rev": rev.rev if rev else None, "rev_status": rev.status.value if rev else None,
            "has_draft": has_draft, "released_rev": released[-1].rev if released else None,
            "tree": _spec_tree(it).model_dump(), "revisions": [{"rev": r.rev, "status": r.status.value, "note": r.note,
                                                                 "released_at": r.released_at, "created_at": r.created_at} for r in it.revisions]}


class LineIn(BaseModel):
    code: str
    name: str = ""
    qty: Decimal = Decimal(1)
    unit: str = "шт"
    item_type: str = "part"


@router.post("/specs/{parent_id}/lines")
def add_line(parent_id: int, data: LineIn, db: Session = Depends(get_db), me: User = Depends(require("bom:write"))):
    """Добавить строку в состав (новую или уже существующую позицию). Пишется в черновую редакцию."""
    parent = db.get(Item, parent_id)
    if not parent:
        raise HTTPException(404)
    code = data.code.strip().upper()
    if not code:
        raise HTTPException(400, "Укажите обозначение")
    child = db.query(Item).filter(Item.code == code).first()
    if not child:
        child = Item(code=code, name=data.name.strip() or code, item_type=ItemType(data.item_type), unit=data.unit, lifecycle=Lifecycle.released)
        db.add(child)
        db.flush()
        db.add(ItemRevision(item_id=child.id, rev="01", status=Lifecycle.released))
    bom_svc.check_cycle(db, parent.id, child.id)
    rev = bom_svc.ensure_draft_revision(db, parent)
    ex = next((l for l in rev.lines if l.child_item_id == child.id), None)
    if ex:
        ex.qty = ex.qty + data.qty
    else:
        db.add(BomLine(parent_revision_id=rev.id, child_item_id=child.id, qty=data.qty, position=len(rev.lines) + 1))
    audit.log(db, me.id, "bom_add", "revision", f"{parent.code}/{rev.rev}", {"child": code, "qty": str(data.qty)})
    db.commit()
    return {"child": _item_brief(child), "rev": rev.rev}


class LinePatch(BaseModel):
    qty: Decimal | None = None
    name: str | None = None
    code: str | None = None
    item_type: str | None = None


@router.patch("/specs/{parent_id}/lines/{line_id}")
def patch_line(parent_id: int, line_id: int, data: LinePatch, db: Session = Depends(get_db), me: User = Depends(require("bom:write"))):
    """Изменить количество (в черновой редакции родителя) и/или переименовать позицию (глобально)."""
    ln = db.get(BomLine, line_id)
    if not ln or ln.parent_revision.item_id != parent_id:
        raise HTTPException(404)
    parent = db.get(Item, parent_id)
    changes = {}
    if data.qty is not None and data.qty != ln.qty:
        rev = bom_svc.ensure_draft_revision(db, parent)
        tgt = next((l for l in rev.lines if l.child_item_id == ln.child_item_id), None)
        if tgt:
            tgt.qty = data.qty
            changes["qty"] = str(data.qty)
    child = ln.child
    if data.name is not None and data.name.strip() and data.name != child.name:
        child.name = data.name.strip()
        changes["name"] = child.name
    if data.code is not None and data.code.strip().upper() != child.code:
        new = data.code.strip().upper()
        if db.query(Item).filter(Item.code == new).first():
            raise HTTPException(400, f"Обозначение {new} уже занято")
        child.code = new
        changes["code"] = new
    if data.item_type and data.item_type != child.item_type.value:
        child.item_type = ItemType(data.item_type)
        changes["item_type"] = data.item_type
    audit.log(db, me.id, "bom_edit", "item", child.code, changes)
    db.commit()
    return {"ok": True, "changes": changes}


@router.delete("/specs/{parent_id}/lines/{line_id}")
def delete_line(parent_id: int, line_id: int, db: Session = Depends(get_db), me: User = Depends(require("bom:write"))):
    ln = db.get(BomLine, line_id)
    if not ln or ln.parent_revision.item_id != parent_id:
        raise HTTPException(404)
    parent = db.get(Item, parent_id)
    rev = bom_svc.ensure_draft_revision(db, parent)
    tgt = next((l for l in rev.lines if l.child_item_id == ln.child_item_id), None)
    if tgt:
        audit.log(db, me.id, "bom_remove", "revision", f"{parent.code}/{rev.rev}", {"child": tgt.child.code})
        db.delete(tgt)
    db.commit()
    return {"ok": True}


@router.post("/specs/{item_id}/apply")
def apply_draft(item_id: int, note: str = "", db: Session = Depends(get_db), me: User = Depends(require("bom:write"))):
    """Применить изменения: черновая редакция становится действующей (у всех затронутых узлов)."""
    it = db.get(Item, item_id)
    if not it:
        raise HTTPException(404)
    applied = []

    def walk(node: Item, seen: set):
        if node.id in seen:
            return
        seen.add(node.id)
        draft = next((r for r in node.revisions if r.status == Lifecycle.draft), None)
        if draft:
            draft.note = draft.note or note or "Правка в редакторе спецификации"
            bom_svc.release_revision(db, draft)
            applied.append(f"{node.code}/{draft.rev}")
        rev = node.current_revision
        for ln in rev.lines if rev else []:
            walk(ln.child, seen)

    # обходим по черновикам, чтобы применить изменения и во вложенных узлах
    def collect(node: Item, seen: set):
        if node.id in seen:
            return
        seen.add(node.id)
        rev = _edit_rev(node)
        for ln in rev.lines if rev else []:
            collect(ln.child, seen)

    order: set = set()
    collect(it, order)
    for node_id in list(order):
        walk(db.get(Item, node_id), set())
    audit.log(db, me.id, "release", "revision", it.code, {"applied": applied})
    db.commit()
    return {"applied": applied}


@router.post("/specs/{item_id}/discard")
def discard_draft(item_id: int, db: Session = Depends(get_db), me: User = Depends(require("bom:write"))):
    it = db.get(Item, item_id)
    if not it:
        raise HTTPException(404)
    n = 0

    def walk(node: Item, seen: set):
        nonlocal n
        if node.id in seen:
            return
        seen.add(node.id)
        for r in list(node.revisions):
            if r.status == Lifecycle.draft and any(x.status == Lifecycle.released for x in node.revisions):
                db.delete(r)
                n += 1
        db.flush()
        rev = node.current_revision
        for ln in rev.lines if rev else []:
            walk(ln.child, seen)

    ids: set = set()

    def collect(node: Item):
        if node.id in ids:
            return
        ids.add(node.id)
        rev = _edit_rev(node)
        for ln in rev.lines if rev else []:
            collect(ln.child)

    collect(it)
    for node_id in list(ids):
        walk(db.get(Item, node_id), set())
    db.commit()
    return {"discarded": n}


@router.post("/import")
def quick_import(file: UploadFile = File(...), db: Session = Depends(get_db), me: User = Depends(require("import:run"))):
    """Импорт в один шаг. Если колонки распознаны — сразу применяем и возвращаем изделие; иначе — предпросмотр."""
    data = file.file.read()
    if len(data) > 30 * 1024 * 1024:
        raise HTTPException(413, "Файл больше 30 МБ")
    try:
        p = excel_import.preview(data)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, f"Не удалось прочитать файл: {e}")
    token = hashlib.sha1(data).hexdigest()[:16]
    _uploads[token] = data
    m = p.mapping
    confident = bool(m.get("code") and (m.get("parent") or m.get("level") or p.indent_detected) and p.rows_total > 0)
    if not confident:
        return {"applied": False, "token": token, "preview": p.model_dump()}
    res = excel_import.apply_bom(db, data, m, p.sheet, p.header_row, None, True)
    audit.log(db, me.id, "import", "bom", token, res, source="import")
    db.commit()
    root = db.query(Item).filter(Item.code == res["roots"][0]).first() if res["roots"] else None
    return {"applied": True, "result": res, "item_id": root.id if root else None, "root_count": len(res["roots"])}


# ------------------------------------------------------------------ заказы --
def _line_status(lines_by_parent: dict[int | None, list[KitLine]], ln: KitLine) -> str:
    if ln.done or ln.issued_qty >= ln.required_qty:
        return "ready"
    kids = lines_by_parent.get(ln.id, [])
    if kids:
        st = [_line_status(lines_by_parent, k) for k in kids]
        if all(s == "ready" for s in st):
            return "partial"  # всё для сборки есть, сам узел ещё не собран
        return "partial" if any(s != "missing" for s in st) else "missing"
    return "partial" if ln.issued_qty > 0 else "missing"


def _kit_summary(kit: Kit) -> dict:
    by_parent: dict[int | None, list[KitLine]] = {}
    for l in kit.lines:
        by_parent.setdefault(l.parent_line_id, []).append(l)
    leaves = [l for l in kit.lines if l.id not in by_parent]
    ready = sum(1 for l in leaves if _line_status(by_parent, l) == "ready")
    today = date.today()
    return {"leaves": len(leaves), "ready": ready, "percent": round(ready * 100 / len(leaves)) if leaves else 100,
            "missing": len(leaves) - ready, "overdue": bool(kit.due_date and kit.due_date < today and kit.status in ("open", "in_work")),
            "days_left": (kit.due_date - today).days if kit.due_date else None}


@router.get("/orders")
def orders(all: bool = False, db: Session = Depends(get_db), me: User = Depends(require("kits:read"))):
    q = db.query(Kit)
    if not all:
        q = q.filter(Kit.status.in_(["open", "in_work"]))
    maxc = 99 if "*" in me.permissions else me.clearance
    out = []
    for k in q.order_by(Kit.due_date.nullslast(), Kit.id.desc()).all():
        if k.item.confidentiality > maxc:
            continue
        pl = db.get(PlanLine, k.plan_line_id) if k.plan_line_id else None
        out.append({"id": k.id, "number": k.number, "item": _item_brief(k.item), "qty": k.qty, "rev": k.revision.rev,
                    "due_date": k.due_date, "status": k.status, "customer": pl.customer if pl else "",
                    "serial_numbers": k.serial_numbers, **_kit_summary(k)})
    return out


class OrderIn(BaseModel):
    item_id: int
    qty: Decimal = Decimal(1)
    due_date: date | None = None
    customer: str = ""
    serial_numbers: str = ""


@router.post("/orders")
def create_order(data: OrderIn, db: Session = Depends(get_db), me: User = Depends(require("kits:write"))):
    it = db.get(Item, data.item_id)
    if not it:
        raise HTTPException(404, "Спецификация не найдена")
    if not it.current_revision or not it.current_revision.lines:
        raise HTTPException(400, "У спецификации пустой состав — сначала заполните его")
    pl = PlanLine(item_id=it.id, qty=data.qty, due_date=data.due_date or date.today(), customer=data.customer, status="launched")
    db.add(pl)
    db.flush()
    kit = plan_svc.create_kit(db, it, data.qty, plan_line_id=pl.id, due_date=data.due_date, serial_numbers=data.serial_numbers)
    audit.log(db, me.id, "create", "kit", kit.number, {"item": it.code, "qty": str(data.qty)})
    db.commit()
    return {"id": kit.id, "number": kit.number}


@router.get("/orders/{kit_id}")
def order(kit_id: int, db: Session = Depends(get_db), me: User = Depends(require("kits:read"))):
    k = db.get(Kit, kit_id)
    if not k:
        raise HTTPException(404, "Заказ не найден")
    check_clearance(me, k.item.confidentiality)
    by_parent: dict[int | None, list[KitLine]] = {}
    for l in k.lines:
        by_parent.setdefault(l.parent_line_id, []).append(l)
    on_order = bom_svc.on_order_map(db)
    stock = bom_svc.stock_map(db)

    def node(l: KitLine) -> dict:
        kids = by_parent.get(l.id, [])
        return {"id": l.id, "item": _item_brief(l.item), "qty_per": l.qty_per, "required": l.required_qty, "received": l.issued_qty,
                "done": l.done, "status": _line_status(by_parent, l), "note": l.note, "level": l.level,
                "on_order": on_order.get(l.item_id, Decimal(0)), "stock": stock.get(l.item_id, Decimal(0)),
                "lead_time_days": l.item.lead_time_days, "children": [node(c) for c in kids]}

    pl = db.get(PlanLine, k.plan_line_id) if k.plan_line_id else None
    return {"id": k.id, "number": k.number, "item": _item_brief(k.item), "qty": k.qty, "rev": k.revision.rev, "due_date": k.due_date,
            "status": k.status, "customer": pl.customer if pl else "", "serial_numbers": k.serial_numbers, "created_at": k.created_at,
            **_kit_summary(k), "tree": [node(l) for l in by_parent.get(None, [])]}


class ReceiveIn(BaseModel):
    received: Decimal | None = None  # явное количество
    done: bool | None = None  # галочка: всё пришло / узел собран


@router.post("/orders/{kit_id}/lines/{line_id}")
def receive(kit_id: int, line_id: int, data: ReceiveIn, db: Session = Depends(get_db), me: User = Depends(require("kits:write"))):
    """Отметка поступления (чек-лист, без складских движений)."""
    k, l = db.get(Kit, kit_id), db.get(KitLine, line_id)
    if not k or not l or l.kit_id != k.id:
        raise HTTPException(404)
    if data.done is not None:
        l.done = data.done
        l.issued_qty = l.required_qty if data.done else Decimal(0)
    if data.received is not None:
        l.issued_qty = max(min(Decimal(data.received), l.required_qty), Decimal(0))
        l.done = l.issued_qty >= l.required_qty
    if k.status == "open":
        k.status = "in_work"
    by_parent: dict[int | None, list[KitLine]] = {}
    for x in k.lines:
        by_parent.setdefault(x.parent_line_id, []).append(x)
    if all(_line_status(by_parent, x) == "ready" for x in by_parent.get(None, [])):
        k.status = "assembled"
    elif k.status == "assembled":
        k.status = "in_work"
    audit.log(db, me.id, "receive", "kit", k.number, {"item": l.item.code, "received": str(l.issued_qty), "done": l.done})
    db.commit()
    return {"received": l.issued_qty, "done": l.done, "kit_status": k.status}


@router.post("/orders/{kit_id}/status")
def order_status(kit_id: int, status: str, db: Session = Depends(get_db), me: User = Depends(require("kits:write"))):
    k = db.get(Kit, kit_id)
    if not k or status not in ("open", "in_work", "assembled", "shipped"):
        raise HTTPException(400)
    k.status = status
    audit.log(db, me.id, "status", "kit", k.number, {"status": status})
    db.commit()
    return {"ok": True}


@router.get("/orders/{kit_id}/export.xlsx")
def export_order(kit_id: int, db: Session = Depends(get_db), me: User = Depends(require("kits:read"))):
    """Ведомость комплектации заказа в Excel (как привыкли плановики)."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    k = db.get(Kit, kit_id)
    if not k:
        raise HTTPException(404)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Комплектация"
    ws.append([f"Заказ {k.number}: {k.item.code} {k.item.name} × {k.qty}", "", "", "", f"срок {k.due_date or '—'}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Уровень", "Обозначение", "Наименование", "Тип", "Нужно", "Ед.", "Пришло", "Не хватает", "Статус", "В заказах у поставщиков"])
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEFC")
    by_parent: dict[int | None, list[KitLine]] = {}
    for l in k.lines:
        by_parent.setdefault(l.parent_line_id, []).append(l)
    on_order = bom_svc.on_order_map(db)
    types = {"assembly": "Сборочная единица", "part": "Деталь", "purchased": "Покупное", "fastener": "Крепёж", "material": "Материал",
             "outsourced": "Кооперация", "outsourced_op": "Операция на стороне", "tooling": "Оснастка", "product": "Изделие"}
    fills = {"ready": "D4EDDA", "partial": "FFF3CD", "missing": "F8D7DA"}

    def walk(l: KitLine):
        st = _line_status(by_parent, l)
        short = max(l.required_qty - l.issued_qty, Decimal(0))
        ws.append([l.level + 1, ("    " * l.level) + l.item.code, l.item.name, types.get(l.item.item_type.value, ""), float(l.required_qty),
                   l.item.unit, float(l.issued_qty), float(short), {"ready": "есть", "partial": "частично", "missing": "нет"}[st],
                   float(on_order.get(l.item_id, 0))])
        for c in ws[ws.max_row]:
            c.fill = PatternFill("solid", fgColor=fills[st])
        for kid in by_parent.get(l.id, []):
            walk(kid)

    for l in by_parent.get(None, []):
        walk(l)
    for col, w in zip("ABCDEFGHIJ", (8, 30, 40, 18, 10, 6, 10, 12, 10, 14)):
        ws.column_dimensions[col].width = w
    ws["B1"].alignment = Alignment(horizontal="left")
    buf = io.BytesIO()
    wb.save(buf)
    from urllib.parse import quote

    name = f"Комплектация {k.number}.xlsx"
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=\"order_{k.id}.xlsx\"; filename*=UTF-8''{quote(name)}"})


# ----------------------------------------------------------------- дефицит --
@router.get("/shortage")
def shortage(db: Session = Depends(get_db), me: User = Depends(require("kits:read"))):
    """Сводный дефицит по открытым заказам: нужно − пришло, листовые позиции."""
    need: dict[int, dict] = {}
    for k in db.query(Kit).filter(Kit.status.in_(["open", "in_work"])).all():
        by_parent: dict[int | None, list[KitLine]] = {}
        for l in k.lines:
            by_parent.setdefault(l.parent_line_id, []).append(l)
        for l in k.lines:
            if l.id in by_parent or l.done:
                continue
            rest = l.required_qty - l.issued_qty
            if rest <= 0:
                continue
            d = need.setdefault(l.item_id, {"item": _item_brief(l.item), "need": Decimal(0), "orders": [], "earliest": None,
                                            "lead_time_days": l.item.lead_time_days, "supplier_id": l.item.default_supplier_id})
            d["need"] += rest
            d["orders"].append({"id": k.id, "number": k.number, "qty": rest, "due_date": k.due_date})
            if k.due_date and (d["earliest"] is None or k.due_date < d["earliest"]):
                d["earliest"] = k.due_date
    on_order = bom_svc.on_order_map(db)
    stock = bom_svc.stock_map(db)
    today = date.today()
    out = []
    for item_id, d in need.items():
        oo = on_order.get(item_id, Decimal(0))
        st = stock.get(item_id, Decimal(0))
        uncovered = max(d["need"] - oo - st, Decimal(0))
        order_by = None
        if d["earliest"]:
            from datetime import timedelta

            order_by = d["earliest"] - timedelta(days=d["lead_time_days"])
        out.append({**d, "need": d["need"].quantize(Q4), "on_order": oo, "stock": st, "uncovered": uncovered.quantize(Q4),
                    "order_by": order_by, "late": bool(order_by and order_by < today and uncovered > 0)})
    out.sort(key=lambda r: (not r["late"], r["order_by"] or date.max, -r["uncovered"]))
    return out


@router.post("/shortage/request")
def shortage_request(db: Session = Depends(get_db), me: User = Depends(require("purchase:write"))):
    """Передать дефицит снабжению: черновики заказов поставщикам (по основному поставщику) + общий без поставщика."""
    from ..models import Partner

    rows = [r for r in shortage(db, me) if r["uncovered"] > 0]
    if not rows:
        return {"orders": []}
    groups: dict[int | None, list] = {}
    for r in rows:
        groups.setdefault(r["supplier_id"], []).append(r)
    created = []
    for supplier_id, items in groups.items():
        if supplier_id is None:
            p = db.query(Partner).filter(Partner.name == "— поставщик не выбран —").first()
            if not p:
                p = Partner(name="— поставщик не выбран —", kind="supplier")
                db.add(p)
                db.flush()
            supplier_id = p.id
        n = db.query(PurchaseOrder).count() + 1
        po = PurchaseOrder(number=f"ЗП-{date.today():%y}-{n:05d}", kind="purchase", partner_id=supplier_id, responsible_id=me.id,
                           due_date=min((i["earliest"] for i in items if i["earliest"]), default=None), comment="Из ведомости дефицита")
        db.add(po)
        db.flush()
        for i in items:
            it = db.get(Item, i["item"]["id"])
            db.add(POLine(order_id=po.id, item_id=it.id, qty=i["uncovered"], price=it.std_cost, due_date=i["earliest"]))
        created.append({"id": po.id, "number": po.number, "lines": len(items)})
    audit.log(db, me.id, "shortage_request", "purchase_order", "", {"orders": created})
    db.commit()
    return {"orders": created}


@router.get("/shortage/export.xlsx")
def export_shortage(db: Session = Depends(get_db), me: User = Depends(require("kits:read"))):
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    rows = shortage(db, me)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Дефицит"
    ws.append(["Обозначение", "Наименование", "Ед.", "Нужно", "На складе", "В заказах", "Не покрыто", "Заказать до", "Для заказов"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEFC")
    for r in rows:
        ws.append([r["item"]["code"], r["item"]["name"], r["item"]["unit"], float(r["need"]), float(r["stock"]), float(r["on_order"]),
                   float(r["uncovered"]), r["order_by"].strftime("%d.%m.%Y") if r["order_by"] else "", ", ".join(o["number"] for o in r["orders"])])
        if r["late"]:
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor="F8D7DA")
    for col, w in zip("ABCDEFGHI", (28, 40, 6, 10, 10, 10, 11, 13, 30)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="deficit.xlsx"'})
